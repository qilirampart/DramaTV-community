from __future__ import annotations

import argparse
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook


UPDATES = {
    7: {
        6: "feed_items 已完成 content_kind + target_type 收口；前台看 content_kind，后端按 target_type 查询真实表，并保留 item_type 兼容旧数据",
        7: "已完成",
        10: "2026-04-21",
    },
    8: {
        6: "已新增 feed_items.content_kind 与 target_type 迁移，并完成发布写入、审核回写、首页查询三段收口；暂未单独拆 content_items 新表",
        7: "已完成",
        10: "2026-04-21",
    },
    20: {
        6: "已落地真实 /me 个人中心：承接个人信息、已发布作品/帖子、获赞、点赞、收藏、草稿箱与资料编辑；私有工作台和公开作者主页已完成拆分",
        7: "已完成",
        10: "2026-04-21",
    },
    36: {
        6: "feed_items 已从旧 item_type 单轴语义升级为 content_kind + target_type 双层语义；item_type 当前仅作为兼容字段保留",
        7: "已完成",
        10: "2026-04-21",
    },
    46: {
        6: "草稿箱已在 /me 落地：接入 draftItems，支持视频/帖子草稿继续编辑与删除；工作流草稿可展示但编辑器仍待开放",
        7: "已完成",
        10: "2026-04-21",
    },
    67: {
        6: "已完成 storage_provider / bucket / object_key / publicBaseUrl 抽象与 URL 解析；真实 OSS 上传、预签名直传和 CDN 仍待运维资源与正式接入",
        7: "开发中",
        10: "2026-04-22",
    },
}


def apply_updates(path: Path, restore_from_backup: bool) -> None:
    backup = Path(f"{path}.bak-20260422")
    original_mode = path.stat().st_mode
    path.chmod(original_mode | 0o200)
    if restore_from_backup and backup.exists():
        try:
            copy2(backup, path)
        finally:
            path.chmod(original_mode | 0o200)

    if not backup.exists():
        copy2(path, backup)

    try:
        wb = load_workbook(path)
        ws = wb.worksheets[0]
        for row_index, changes in UPDATES.items():
            for col_index, value in changes.items():
                ws.cell(row_index, col_index).value = value
        wb.save(path)
    finally:
        path.chmod(original_mode)
    print(f"updated: {path}")
    print(f"backup: {backup}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("paths", nargs="+")
    parser.add_argument("--restore-from-backup", action="store_true")
    args = parser.parse_args()

    for raw_path in args.paths:
        apply_updates(Path(raw_path), restore_from_backup=args.restore_from_backup)


if __name__ == "__main__":
    main()
